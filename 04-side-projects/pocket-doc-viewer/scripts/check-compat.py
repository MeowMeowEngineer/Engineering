"""앱에서 저장한 파일이 실제 프로그램 형식 규칙에 맞는지 검사한다.

python3 scripts/check-compat.py <폴더>
- HWPX: python-hwpx로 다시 열기·패키지 검사, mimetype이 첫 항목이고 무압축인지
- DOCX/PPTX/XLSX: python-docx·python-pptx·openpyxl로 열기, 모든 XML 파트가 올바른 XML인지(선언 1개)
- PDF: pypdf로 열기
- LibreOffice가 있으면 DOCX·PPTX·XLSX를 PDF로 변환해 실제로 열리는지 확인
"""
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from xml.dom import minidom

folder = Path(sys.argv[1] if len(sys.argv) > 1 else "edit-out")
failures = []


def check(name, ok, detail=""):
    print(("PASS " if ok else "FAIL ") + name + (f" - {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def xml_parts_ok(path):
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if not info.filename.endswith((".xml", ".rels", ".hpf")):
                continue
            data = z.read(info)
            if data.count(b"<?xml") > 1:
                return False, f"{info.filename}: XML 선언 중복"
            try:
                minidom.parseString(data)
            except Exception as e:  # noqa: BLE001
                return False, f"{info.filename}: {e}"
    return True, ""


for f in sorted(folder.iterdir()):
    ext = f.suffix.lower()
    if ext in (".docx", ".pptx", ".xlsx", ".hwpx"):
        ok, detail = xml_parts_ok(f)
        check(f"{f.name} XML 파트", ok, detail)
    if ext == ".hwpx":
        from hwpx import HwpxDocument, validate_editor_open_safety

        z = zipfile.ZipFile(f)
        first = z.infolist()[0]
        check(f"{f.name} mimetype 첫 항목·무압축", first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED)
        report = validate_editor_open_safety(str(f))
        check(f"{f.name} 다시 열기", report.reopen_ok and not report.blocking_package_errors, str(report.reopen_error or ""))
        HwpxDocument.open(str(f))
    elif ext == ".docx":
        import docx

        check(f"{f.name} python-docx", len(docx.Document(str(f)).paragraphs) > 0)
    elif ext == ".pptx":
        from pptx import Presentation

        check(f"{f.name} python-pptx", len(Presentation(str(f)).slides) > 0)
    elif ext == ".xlsx":
        import openpyxl

        check(f"{f.name} openpyxl", len(openpyxl.load_workbook(str(f)).sheetnames) > 0)
    elif ext == ".pdf":
        import pypdf

        check(f"{f.name} pypdf", len(pypdf.PdfReader(str(f)).pages) > 0)

soffice = shutil.which("soffice")
if soffice:
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() not in (".docx", ".pptx", ".xlsx"):
            continue
        with tempfile.TemporaryDirectory() as out:
            subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", out, str(f)],
                           capture_output=True, timeout=180)
            check(f"{f.name} LibreOffice 열기", any(Path(out).glob("*.pdf")))
else:
    print("SKIP LibreOffice 없음")

print("모두 통과" if not failures else f"실패 {len(failures)}건")
sys.exit(1 if failures else 0)
