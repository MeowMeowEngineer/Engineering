"""데모와 테스트에 쓰는 샘플 문서를 만든다.

python3 scripts/make-samples.py  ->  public/samples/
필요 패키지: python-docx openpyxl python-pptx reportlab python-hwpx pillow
"""
from pathlib import Path
import io

from PIL import Image, ImageDraw

OUT = Path(__file__).resolve().parent.parent / "public" / "samples"
OUT.mkdir(parents=True, exist_ok=True)

QUARTERS = ["1분기", "2분기", "3분기", "4분기"]
SALES = {"서울": [120, 135, 150, 170], "부산": [80, 92, 88, 101], "대구": [55, 60, 72, 69]}


def chart_png() -> bytes:
    img = Image.new("RGB", (480, 260), "white")
    d = ImageDraw.Draw(img)
    colors = ["#2f6fde", "#e8833a", "#3aa76d"]
    for ci, (city, vals) in enumerate(SALES.items()):
        for qi, v in enumerate(vals):
            x = 40 + qi * 110 + ci * 28
            d.rectangle([x, 230 - v, x + 22, 230], fill=colors[ci])
    d.line([30, 230, 470, 230], fill="#333", width=2)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def make_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    doc.add_heading("포켓 문서 뷰어 샘플 보고서", 0)
    p = doc.add_paragraph("이 문서는 ")
    p.add_run("Word(.docx)").bold = True
    p.add_run(" 렌더링을 확인하기 위한 샘플입니다. ")
    r = p.add_run("굵게, 기울임, 색상")
    r.italic = True
    r.font.color.rgb = RGBColor(0x2F, 0x6F, 0xDE)
    p.add_run("이 모두 표시되어야 합니다.")
    doc.add_heading("1. 분기별 매출", 1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["지역", *QUARTERS]):
        t.rows[0].cells[i].text = h
    for city, vals in SALES.items():
        row = t.add_row().cells
        row[0].text = city
        for i, v in enumerate(vals):
            row[i + 1].text = f"{v:,}"
    doc.add_heading("2. 차트", 1)
    doc.add_picture(io.BytesIO(chart_png()), width=Inches(4.5))
    doc.add_heading("3. 정리", 1)
    for s in ["모든 문서는 기기 안에서만 열립니다.", "광고와 회원가입이 없습니다.", "인터넷 연결 없이도 동작합니다."]:
        doc.add_paragraph(s, style="List Bullet")
    for i in range(6):
        para = doc.add_paragraph(
            "문단 %d. 휴대폰 화면 폭에 맞춰 페이지 전체가 축소되어 보이고, 두 손가락으로 확대할 수 있습니다." % (i + 1)
        )
        para.runs[0].font.size = Pt(11)
    doc.save(OUT / "sample.docx")


def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "매출"
    ws.merge_cells("A1:F1")
    ws["A1"] = "2026년 지역별 분기 매출 (단위: 백만 원)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(["지역", *QUARTERS, "합계"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6FDE")
    for city, vals in SALES.items():
        ws.append([city, *vals, sum(vals)])
    totals = [sum(v[i] for v in SALES.values()) for i in range(4)]
    ws.append(["합계", *totals, sum(totals)])
    ws.column_dimensions["A"].width = 12
    ws2 = wb.create_sheet("직원")
    ws2.append(["이름", "부서", "입사일", "연차"])
    names = ["김민준", "이서연", "박도윤", "최하은", "정시우", "강지우", "조예준", "윤수아"]
    for i, n in enumerate(names):
        ws2.append([n, ["개발", "영업", "디자인"][i % 3], f"2020-0{i % 9 + 1}-15", i + 1])
    wb.save(OUT / "sample.xlsx")


def make_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "포켓 문서 뷰어"
    s.placeholders[1].text = "휴대폰에서 여는 무료 문서 뷰어\nWord · Excel · PowerPoint · PDF · 한글"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "지원 형식"
    body = s.placeholders[1].text_frame
    body.text = "PDF: pdf.js로 페이지 렌더링"
    for line in ["Word(.docx): docx-preview", "Excel(.xlsx, .xls, .csv): SheetJS", "PowerPoint(.pptx): 직접 만든 렌더러", "한글(.hwp, .hwpx): hwp.js + 직접 만든 HWPX 렌더러"]:
        body.add_paragraph().text = line

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "도형과 그림"
    colors = [(0x2F, 0x6F, 0xDE), (0xE8, 0x83, 0x3A), (0x3A, 0xA7, 0x6D)]
    for i, label in enumerate(["열기", "보기", "확대"]):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 2.4), Inches(2.2), Inches(2.0), Inches(1.2))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*colors[i])
        shp.line.fill.background()
        tf = shp.text_frame
        tf.text = label
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.paragraphs[0].runs[0].font.size = Pt(28)
        tf.paragraphs[0].runs[0].font.bold = True
        tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    s.shapes.add_picture(io.BytesIO(chart_png()), Inches(8.2), Inches(1.8), width=Inches(4.6))
    tb = s.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(6.8), Inches(2))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "도형 채우기, 글자 색, 그림 위치를 슬라이드 좌표 그대로 그립니다."
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "표"
    rows, cols = len(SALES) + 1, 5
    tbl = s.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(3)).table
    for i, h in enumerate(["지역", *QUARTERS]):
        tbl.cell(0, i).text = h
    for r, (city, vals) in enumerate(SALES.items(), start=1):
        tbl.cell(r, 0).text = city
        for c, v in enumerate(vals, start=1):
            tbl.cell(r, c).text = str(v)
    prs.save(OUT / "sample.pptx")


def make_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    c = canvas.Canvas(str(OUT / "sample.pdf"), pagesize=A4)
    w, h = A4
    for page in range(1, 4):
        c.setFillColorRGB(0.18, 0.44, 0.87)
        c.rect(0, h - 90, w, 90, fill=1, stroke=0)
        c.setFillColorRGB(1, 1, 1)
        c.setFont("HYGothic-Medium", 24)
        c.drawString(50, h - 58, f"포켓 문서 뷰어 PDF 샘플 · {page}쪽")
        c.setFillColorRGB(0, 0, 0)
        c.setFont("HYSMyeongJo-Medium", 13)
        y = h - 140
        for i in range(24):
            c.drawString(50, y, f"{i + 1:02d}. 한글 CID 글꼴(CMap)이 필요한 PDF도 깨지지 않고 표시되어야 합니다.")
            y -= 26
        c.setFont("Helvetica", 10)
        c.drawCentredString(w / 2, 30, f"- {page} -")
        c.showPage()
    c.save()


def make_hwpx():
    from hwpx import HwpxDocument

    doc = HwpxDocument.new()
    doc.add_paragraph("포켓 문서 뷰어 HWPX 샘플")
    doc.add_paragraph("이 문서는 한글(.hwpx) 렌더링을 확인하기 위한 샘플입니다. 문단, 표, 글자 모양을 직접 해석해 그립니다.")
    doc.add_paragraph("")
    doc.add_paragraph("1. 분기별 매출")
    tbl = doc.add_table(len(SALES) + 1, 5)
    for i, h in enumerate(["지역", *QUARTERS]):
        tbl.set_cell_text(0, i, h)
    for r, (city, vals) in enumerate(SALES.items(), start=1):
        tbl.set_cell_text(r, 0, city)
        for c, v in enumerate(vals, start=1):
            tbl.set_cell_text(r, c, str(v))
    doc.add_paragraph("")
    doc.add_paragraph("2. 특징")
    for s in ["기기 안에서만 문서를 엽니다.", "광고와 회원가입이 없습니다.", "공공기관 서식처럼 표가 많은 문서도 표시합니다."]:
        doc.add_paragraph("  · " + s)
    doc.save_to_path(str(OUT / "sample.hwpx"))
    (OUT / "sample.hwpx").chmod(0o644)


def make_csv():
    lines = ["지역," + ",".join(QUARTERS)] + [f"{c}," + ",".join(map(str, v)) for c, v in SALES.items()]
    (OUT / "sample.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    for fn in (make_docx, make_xlsx, make_pptx, make_pdf, make_hwpx, make_csv):
        fn()
        print("ok", fn.__name__)
