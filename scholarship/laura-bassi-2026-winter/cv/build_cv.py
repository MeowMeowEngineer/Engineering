"""Build Eunsung_Jang_CV.docx (A4) for the Laura Bassi application.

Run: python3 build_cv.py, then
     soffice --headless --convert-to pdf Eunsung_Jang_CV.docx
Text in [CONFIRM: ...] is highlighted yellow and must be filled or removed before upload.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("Eunsung_Jang_CV.docx")
FONT = "Calibri"
BODY_PT = 10
TEXT_WIDTH_CM = 21.0 - 2 * 1.8
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

# **bold** and [CONFIRM: ...] markup inside any text string
TOKEN = re.compile(r"(\*\*.+?\*\*|\[CONFIRM:[^\]]*\])")


def runs(par, text, size=BODY_PT, italic=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        bold = part.startswith("**")
        r = par.add_run(part[2:-2] if bold else part)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if part.startswith("[CONFIRM"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


def spacing(par, before=0, after=2, line=1.08):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def heading(doc, title):
    p = doc.add_paragraph()
    spacing(p, before=9, after=3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "1F3A5F"}.items():
        bottom.set(qn(k), v)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def dated(doc, left, right, sub=None):
    """Bold entry with a right-aligned date, optional second line."""
    p = doc.add_paragraph()
    spacing(p, before=2, after=0 if sub else 2)
    p.paragraph_format.keep_with_next = bool(sub)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    runs(p, f"**{left}**\t")
    runs(p, right)
    if sub:
        q = doc.add_paragraph()
        spacing(q, after=2)
        runs(q, sub, italic=False)


def bullet(doc, text, label=None):
    p = doc.add_paragraph()
    spacing(p, after=2)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.9)
    pf.first_line_indent = Cm(-0.9)
    pf.tab_stops.add_tab_stop(Cm(0.9))
    runs(p, f"{label or '•'}\t{text}")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    p = doc.add_paragraph()
    spacing(p, after=0)
    r = p.add_run("Eunsung Jang")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    spacing(p, after=0)
    runs(p, "Integrated M.S./Ph.D. student, RF System Technologies Lab., Department of Electrical and "
            "Computer Engineering, Ajou University, Suwon, Republic of Korea")
    p = doc.add_paragraph()
    spacing(p, after=4)
    runs(p, "Email: jangstar@ajou.ac.kr")

    heading(doc, "Research Interests")
    for t in (
        "Dual-polarized distributed FMCW radar with lens array antennas for through-the-wall human detection",
        "AI radar signal processing for through-the-wall human detection",
        "Wide-angle scanning phased array antenna systems for anti-drone detection",
        "Low-loss coaxial cables",
    ):
        bullet(doc, t)

    heading(doc, "Education")
    dated(doc, "Ajou University, Suwon, Korea", "[CONFIRM: start year]-present",
          "Integrated M.S./Ph.D. program, Department of Electrical and Computer Engineering. "
          "Advisor: Prof. Jeong-Wook Kim")
    dated(doc, "[CONFIRM: university]", "[CONFIRM: year]",
          "B.S. in [CONFIRM: major]")

    heading(doc, "Professional Experience")
    dated(doc, "Hanwha Ocean, Shipbuilding Division", "2015-2017", "HVAC Design Engineer")

    heading(doc, "Publications")
    sub = doc.add_paragraph()
    spacing(sub, before=2, after=1)
    sub.paragraph_format.keep_with_next = True
    runs(sub, "**Journal articles**")
    bullet(doc, "**E. Jang**, J.-W. Park, K.-H. Park, J.-Y. Lee, H.-S. Kim, and J.-W. Kim, \"Design of "
                "dual-polarized distributed FMCW radar system with high-gain lens array antenna for "
                "through-the-wall human detection application,\" IEEE Internet of Things Journal, 2025.", "[J1]")
    bullet(doc, "**E. Jang**, E. You, and J.-W. Kim, \"Coaxial cable loss reduction through uniform air-hole "
                "PTFE dielectric fabricated by ram extrusion,\" under review.", "[J2]")
    sub = doc.add_paragraph()
    spacing(sub, before=4, after=1)
    sub.paragraph_format.keep_with_next = True
    runs(sub, "**International conference**")
    bullet(doc, "**E. Jang**, J.-W. Park, K. Park, H. Kim, J. Lee, and J.-W. Kim, \"Machine learning selection "
                "of range bins for coherent integration in distributed radar observing frontal and lateral "
                "motion,\" in Proc. Asia-Pacific Microwave Conference (APMC), 2025.", "[C1]")
    sub = doc.add_paragraph()
    spacing(sub, before=4, after=1)
    sub.paragraph_format.keep_with_next = True
    runs(sub, "**Domestic conferences (in Korean)**")
    bullet(doc, "**E. Jang**, S. Choi, M. Cho, J.-W. Park, S. Kim, and J.-W. Kim, \"Design of a dual-polarized "
                "high-gain lens-coupled stacked patch antenna unit cell for C-band ISAC applications,\" "
                "KICS Summer Conference, 2026.", "[D1]")
    bullet(doc, "**E. Jang**, J.-W. Park, K.-H. Park, J.-Y. Lee, H.-S. Kim, and J.-W. Kim, \"Through-the-wall "
                "human detection using band-energy-difference-based channel combining in a distributed "
                "dual-polarized FMCW radar,\" KIEES Winter Conference, 2026.", "[D2]")
    bullet(doc, "**E. Jang**, J.-W. Park, J.-W. Kim, K.-H. Park, J.-Y. Lee, and H.-S. Kim, \"Dual-polarized "
                "distributed FMCW radar for through-the-wall human detection,\" KIEES Summer Conference, "
                "2025.", "[D3]")

    heading(doc, "Awards and Scholarships")
    for year, text in (
        ("2026", "Master's Excellence Scholarship (Science and Engineering), [CONFIRM: awarding body]"),
        ("2025", "Gold Prize, BK21 Graduate School Innovation Outstanding Paper Award "
                 "[CONFIRM: 2025 or 2026]"),
        ("2025", "Encouragement Award, 24th Graduate Student Paper Contest in Radio Engineering, for \"Design "
                 "of a dual-polarized distributed FMCW radar system with a high-gain lens array antenna and "
                 "its application to through-the-wall human detection\""),
        ("2024", "Encouragement Award, Industry-Academia Fair, Department of Electrical and Computer "
                 "Engineering, Ajou University"),
    ):
        bullet(doc, text, year)

    heading(doc, "Research Experience")
    for title, text in (
        ("Through-the-wall human detection radar.",
         "Designed a dual-polarized distributed FMCW radar with high-gain lens array antennas; developed "
         "machine learning range-bin selection and polarization channel combining [J1, C1, D2, D3]."),
        ("Wide-angle dual-polarized phased array for drone detection.",
         "X-band, ±80° beam scanning, LHCP/RHCP. Carried out design, fabrication, measurement, test-jig "
         "design, chamber measurement at RAPA, and beamforming experiments on my own."),
        ("C-band ISAC antenna.",
         "Dual-polarized high-gain lens-coupled stacked patch antenna unit cell [D1]."),
        ("Low-loss coaxial cable.",
         "EM simulation, design optimization, manuscript writing, and response to reviewers [J2]."),
        ("X-band analog beamforming SAR.",
         "Wrote the ADAR1000 SPI beam-control manual and a debugging protocol for reverse voltage in "
         "the receiver."),
        ("Antenna design automation.",
         "Built a workflow in which a large language model drives a 3D EM simulator from text "
         "instructions; modeling time fell by more than 95%, and the loop ran over 1,000 simulations "
         "in 5 days."),
        ("Research proposals.",
         "Designed the structure of and coordinated five national R&D proposals, each written in 1-2 days."),
    ):
        bullet(doc, f"**{title}** {text}")

    heading(doc, "Skills")
    for label, text in (
        ("EM design", "3D full-wave EM simulation, Keysight ADS, openEMS, MATLAB"),
        ("RF measurement", "VNA, spectrum analyzer, oscilloscope, signal generator, power meter, "
                           "anechoic-chamber antenna measurement"),
        ("Hardware", "PCB design (PADS), soldering, breadboard prototyping, test-jig design, "
                     "ADAR1000 control over SPI"),
        ("Programming", "Python, MATLAB, JavaScript, Java (Android), VBA"),
        ("Languages", "Korean (native), English [CONFIRM: test score, if any]"),
    ):
        bullet(doc, f"**{label}:** {text}")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
